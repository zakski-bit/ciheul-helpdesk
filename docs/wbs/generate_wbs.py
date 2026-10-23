import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def build_wbs_spreadsheet():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "WBS Helpdesk System"
    ws.views.sheetView[0].showGridLines = True

    # ----------------------------------------------------
    # Professional Corporate Palette (Clean, Simple, Elegant)
    # ----------------------------------------------------
    NAVY_PRIMARY = "1E293B"      # Deep Slate Navy for Main Banner & Header
    SLATE_HEADER = "334155"      # Subtle Slate for Table Header
    SECTION_FILL = "F1F5F9"      # Very soft cool grey for Module Section rows
    ZEBRA_FILL = "F8FAFC"        # Very subtle alternating row
    WHITE = "FFFFFF"
    BORDER_SUBTLE = "CBD5E1"     # Thin refined slate border
    BORDER_MEDIUM = "64748B"

    # Fonts
    font_main_title = Font(name="Segoe UI", size=13, bold=True, color=WHITE)
    font_subtitle = Font(name="Segoe UI", size=9.5, italic=False, color="E2E8F0")
    font_header = Font(name="Segoe UI", size=10, bold=True, color=WHITE)
    font_section = Font(name="Segoe UI", size=9.5, bold=True, color="0F172A")
    font_body = Font(name="Segoe UI", size=9.5, color="1E293B")
    font_bold_body = Font(name="Segoe UI", size=9.5, bold=True, color="0F172A")

    # Fills
    fill_main_title = PatternFill(start_color=NAVY_PRIMARY, end_color=NAVY_PRIMARY, fill_type="solid")
    fill_header = PatternFill(start_color=SLATE_HEADER, end_color=SLATE_HEADER, fill_type="solid")
    fill_section = PatternFill(start_color=SECTION_FILL, end_color=SECTION_FILL, fill_type="solid")
    fill_zebra = PatternFill(start_color=ZEBRA_FILL, end_color=ZEBRA_FILL, fill_type="solid")
    fill_white = PatternFill(start_color=WHITE, end_color=WHITE, fill_type="solid")
    fill_summary = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

    # Borders
    thin_border = Border(
        left=Side(style="thin", color=BORDER_SUBTLE),
        right=Side(style="thin", color=BORDER_SUBTLE),
        top=Side(style="thin", color=BORDER_SUBTLE),
        bottom=Side(style="thin", color=BORDER_SUBTLE)
    )
    summary_border = Border(
        left=Side(style="thin", color=BORDER_SUBTLE),
        right=Side(style="thin", color=BORDER_SUBTLE),
        top=Side(style="thin", color=BORDER_MEDIUM),
        bottom=Side(style="double", color=NAVY_PRIMARY)
    )

    # Alignments
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_right = Alignment(horizontal="right", vertical="center")

    # 1. Header Block (Simple & Corporate)
    ws.merge_cells("A1:J1")
    ws["A1"] = "CIHEUL TECHNOLOGIES - TAKE HOME PROGRAMMING TEST"
    ws["A1"].font = font_main_title
    ws["A1"].fill = fill_main_title
    ws["A1"].alignment = align_center

    ws.merge_cells("A2:J2")
    ws["A2"] = "Work Breakdown Structure (WBS) - Omnichannel Helpdesk & Real-time Support System"
    ws["A2"].font = Font(name="Segoe UI", size=10.5, bold=True, color=WHITE)
    ws["A2"].fill = fill_main_title
    ws["A2"].alignment = align_center

    ws.merge_cells("A3:J3")
    ws["A3"] = "Arsitektur: Monorepo | Backend: Python FastAPI | Frontend: Vanilla JS Widget & Tailwind Web Dashboard | Mobile: Flutter"
    ws["A3"].font = font_subtitle
    ws["A3"].fill = fill_main_title
    ws["A3"].alignment = align_center

    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 18
    ws.row_dimensions[4].height = 8  # Subtle spacer row

    # 2. Table Column Headers
    headers = [
        ("WBS ID", 10),
        ("Modul / Komponen", 22),
        ("Fitur / Work Package", 28),
        ("Rincian Task & Deliverables Teknis", 46),
        ("User Role", 16),
        ("Tech Stack", 22),
        ("Durasi (Hari)", 13),
        ("Bobot (%)", 11),
        ("Prioritas", 13),
        ("Kriteria Verifikasi (Definition of Done)", 42)
    ]

    header_row = 5
    ws.row_dimensions[header_row].height = 26
    for col_idx, (header_text, width) in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=header_text)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = thin_border
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    # 3. WBS Items
    items = [
        # 1.0 Analisis & Desain Sistem
        (True, "1.0", "Analisis & Desain", "Spesifikasi & Blueprint Sistem", "Analisis kebutuhan teknis dan proses bisnis helpdesk omnichannel", "Semua Role", "Markdown, Draw.io", 2.0, 8.0, "Critical", "Dokumen spesifikasi terverifikasi"),
        (False, "1.1", "Analisis & Desain", "BPMN Swimlane Diagram", "Pemodelan proses bisnis 4 swimlanes (Customer, AI Bot, CS Agent, Supervisor)", "Lead Architect", "Draw.io / XML", 1.0, 4.0, "Critical", "File .drawio dan export PDF bersih"),
        (False, "1.2", "Analisis & Desain", "WBS & System Design", "Penyusunan rincian WBS modular dan dokumen arsitektur database (ERD) & API contract", "Tech Lead", "Excel / Markdown", 1.0, 4.0, "Critical", "File .xlsx terformat rapi dan Architecture.md"),

        # 2.0 Backend Core & Database Engine
        (True, "2.0", "Backend Service", "Core API & Database Engine", "Fondasi RESTful API, database schema, otentikasi JWT dan manajemen tiket", "CS, Admin", "Python FastAPI", 3.0, 15.0, "Critical", "Backend running, API docs /docs aktif"),
        (False, "2.1", "Backend Service", "Database Schema & ORM", "Modeling tabel Users, Customers, Tickets, Messages, KnowledgeBase dengan relasi ForeignKey", "Backend Eng", "SQLite / SQLAlchemy", 1.0, 5.0, "Critical", "Migrasi schema sukses, relasi cascade teruji"),
        (False, "2.2", "Backend Service", "Auth & RBAC Service", "Autentikasi agen CS dan Admin berbasis JWT Token, password hash bcrypt", "Admin, CS Agent", "FastAPI Security, JWT", 1.0, 5.0, "High", "Endpoint login, refresh token, validasi role"),
        (False, "2.3", "Backend Service", "Ticket Management CRUD", "API pembuatan tiket baru, update status (OPEN, IN_PROGRESS, RESOLVED), filter kategori", "CS Agent, Customer", "FastAPI, Pydantic", 1.0, 5.0, "Critical", "CRUD tiket berfungsi dengan response DTO tervalidasi"),

        # 3.0 Real-time WebSocket Hub (Always-Connected)
        (True, "3.0", "Real-time Hub", "WebSocket Infrastructure", "Jalur komunikasi real-time 3-arah: Live Chat Widget <-> Web Dashboard <-> Mobile Flutter", "Customer, CS", "FastAPI WebSockets", 3.0, 16.0, "Critical", "Pesan tersinkronisasi instan <100ms"),
        (False, "3.1", "Real-time Hub", "Connection Manager & Rooms", "Manajemen pool koneksi aktif, perutean broadcast pesan berdasarkan ticket_id session", "Backend Eng", "Python asyncio", 1.0, 6.0, "Critical", "Multi-client join room tanpa memory leak"),
        (False, "3.2", "Real-time Hub", "Always-Connected & Heartbeat", "Mekanisme ping-pong 15s deteksi zombie socket, reconnect otomatis", "Backend & Frontend", "WebSocket Protocol", 1.0, 5.0, "Critical", "Koneksi otomatis pulih saat jaringan putus-nyambung"),
        (False, "3.3", "Real-time Hub", "Offline Message Catch-up", "Sinkronisasi pesan tertinggal dengan parameter last_received_message_id saat reconnect", "Backend Eng", "SQLAlchemy + WS", 1.0, 5.0, "High", "Tidak ada pesan yang hilang saat client offline"),

        # 4.0 AI Bot Assistant & Escalation Engine
        (True, "4.0", "AI Bot Engine", "Auto-Reply & Handoff Triage", "Penanganan otomatis keluhan awal pelanggan dan eskalasi mulus ke CS manusia", "Customer, AI Bot", "Python, NLP / LLM API", 2.5, 13.0, "High", "Bot menjawab akurat dan eskalasi otomatis"),
        (False, "4.1", "AI Bot Engine", "Knowledge Base Matching", "Pencocokan pertanyaan keluhan dengan FAQ Knowledge Base menggunakan similarity scoring", "AI Bot", "Python, TF-IDF / Gemini", 1.0, 5.0, "High", "Rekomendasi jawaban bot relevan (score >= 0.75)"),
        (False, "4.2", "AI Bot Engine", "Confidence Triage & Fallback", "Deteksi jawaban tidak meyakinkan atau sentimen komplain untuk memicu eskalasi", "AI Bot", "State Machine Logic", 0.5, 3.0, "High", "Status tiket berganti WAITING_AGENT secara otomatis"),
        (False, "4.3", "AI Bot Engine", "Human Handoff Workflow", "Broadcast event ticket_escalated ke Web/Mobile CS dan nonaktifkan bot saat di-claim", "CS Agent, Bot", "WebSocket Hub", 1.0, 5.0, "Critical", "Transisi dari bot ke CS mulus tanpa duplikasi chat"),

        # 5.0 Live Chat Widget (JavaScript)
        (True, "5.0", "Live Chat Widget", "Embeddable Client Script", "Widget chat melayang di kanan bawah website (ala Intercom) dengan Vanilla JS murni", "Customer", "Vanilla JS, Scoped CSS", 2.5, 12.0, "Critical", "Widget tampil responsif, bundle < 40 KB"),
        (False, "5.1", "Live Chat Widget", "Floating Launcher & Modal", "Tombol trigger mengambang (FAB) dan popup window chat responsif tanpa merusak CSS host", "Frontend Eng", "HTML5, Scoped CSS", 1.0, 4.0, "Critical", "UI responsif di mobile & desktop, bebas conflict"),
        (False, "5.2", "Live Chat Widget", "WebSocket Client & UI Feed", "Koneksi WebSocket client, auto-reconnect, pengiriman pesan, bubbles chat, typing indicator", "Frontend Eng", "Native WebSocket API", 1.0, 5.0, "Critical", "Pesan masuk real-time, status terkirim jelas"),
        (False, "5.3", "Live Chat Widget", "Interactive Action Buttons", "Tombol eskalasi 'Hubungi CS Manusia', feedback rating kepuasan setelah masalah selesai", "Frontend Eng", "Vanilla JS Events", 0.5, 3.0, "High", "User dapat memicu eskalasi langsung lewat klik"),

        # 6.0 Web App CS & Admin Dashboard
        (True, "6.0", "Web Dashboard", "CS Workspace & Admin Panel", "Aplikasi web internal untuk agen CS melayani chat dan admin memantau metrik performa", "CS Agent, Admin", "HTML5, Tailwind CSS, JS", 3.0, 14.0, "Critical", "Dashboard rapi, realtime feed, clean UI"),
        (False, "6.1", "Web Dashboard", "Clean Responsive UI", "Desain modern, kontras WCAG AA, bebas indikator klise", "UI/UX & Frontend", "Tailwind CSS", 1.0, 4.0, "High", "Tampilan modern, profesional"),
        (False, "6.2", "Web Dashboard", "Real-time Ticket Kanban", "Board tiket aktif dengan badge status (BOT, WAITING, ACTIVE, RESOLVED), claim button", "Frontend Eng", "Modular JS, WS", 1.0, 5.0, "Critical", "Tiket baru muncul seketika tanpa refresh browser"),
        (False, "6.3", "Web Dashboard", "Interactive 1-on-1 Chat Room", "Panel percakapan langsung dengan customer, informasi identitas user, ganti status tiket", "Frontend Eng", "WebSocket Client", 1.0, 5.0, "Critical", "CS dapat membalas chat real-time dan resolve tiket"),

        # 7.0 Mobile App CS (Flutter)
        (True, "7.0", "Mobile App (CS)", "Flutter Android Client", "Aplikasi mobile resmi agen CS untuk melayani tiket kendala di mana saja", "CS Agent", "Flutter (Dart)", 3.0, 12.0, "Critical", "Aplikasi Flutter jalan lancar, siap build APK"),
        (False, "7.1", "Mobile App (CS)", "Architecture & State Setup", "Scaffolding project Flutter, routing, state management untuk list tiket dan chat room", "Mobile Eng", "Flutter, Provider", 1.0, 4.0, "Critical", "Clean architecture, reactive state management"),
        (False, "7.2", "Mobile App (CS)", "Real-time WebSocket Client", "Implementasi client WebSocket di Flutter dengan background awareness & auto-reconnect", "Mobile Eng", "web_socket_channel", 1.0, 4.0, "Critical", "Pesan sync real-time antara Mobile dan Web"),
        (False, "7.3", "Mobile App (CS)", "In-App Notifications & Resolution", "Pemberitahuan tiket eskalasi baru dengan banner alert dan unread badge counter", "Mobile Eng", "Flutter Widgets", 1.0, 4.0, "High", "Badge notifikasi dan alert tiket baru aktif"),

        # 8.0 QA, Security Audit & Automated Release (CI/CD)
        (True, "8.0", "QA & Release", "Verification & Delivery", "Pengujian menyeluruh, audit keamanan kode, dan otomasi build APK via GitHub Actions", "QA, DevOps", "Pytest, GitHub Actions", 2.0, 10.0, "Critical", "Test lulus 100%, APK ter-compile otomatis"),
        (False, "8.1", "QA & Release", "Unit & Security Testing", "Test suite pytest (API, WebSocket, AI deflection), sanitasi XSS pada input chat widget", "QA Eng", "Pytest, TestClient", 1.0, 4.0, "Critical", "Zero vulnerabilities, seluruh unit test hijau"),
        (False, "8.2", "QA & Release", "GitHub Actions CI (Build APK)", "Workflow CI/CD untuk otomatis compile Flutter release APK pada commit/tag repository", "DevOps Eng", "GitHub Actions YAML", 0.5, 3.0, "Critical", "APK file artifact terbuat otomatis di GitHub"),
        (False, "8.3", "QA & Release", "Documentation & Verification", "Penyusunan README lengkap, panduan run lokal 1-command, verifikasi link submission", "Tech Lead", "Markdown, Git", 0.5, 3.0, "Critical", "Seluruh deliverable take-home test siap kirim")
    ]

    current_row = header_row + 1
    for item in items:
        is_section, wbs_id, modul, fitur, task, role, tech, durasi, bobot, prioritas, dod = item
        # Do not force hardcoded small row heights; allow Excel AutoFit to expand rows
        ws.row_dimensions[current_row].height = None

        vals = [wbs_id, modul, fitur, task, role, tech, durasi, bobot, prioritas, dod]
        for col_idx, val in enumerate(vals, start=1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.border = thin_border

            if is_section:
                cell.fill = fill_section
                cell.font = font_section
            else:
                cell.font = font_body
                if current_row % 2 == 0:
                    cell.fill = fill_zebra
                else:
                    cell.fill = fill_white

            # Alignments & formatting
            if col_idx in [1, 5, 6, 9]:
                cell.alignment = align_center
            elif col_idx == 7:  # Durasi
                cell.alignment = align_right
                cell.number_format = "0.0"
            elif col_idx == 8:  # Bobot
                cell.alignment = align_right
                cell.number_format = '0.0"%"'
            else:
                cell.alignment = align_left

        current_row += 1

    # 4. Total Summary Row
    summary_row = current_row
    ws.row_dimensions[summary_row].height = 24
    ws.merge_cells(f"A{summary_row}:F{summary_row}")
    summary_label = ws.cell(row=summary_row, column=1, value="TOTAL ESTIMASI PROYEK (HARI KERJA & BOBOT)")
    summary_label.font = font_bold_body
    summary_label.fill = fill_summary
    summary_label.alignment = align_right

    for c in range(1, 7):
        ws.cell(row=summary_row, column=c).border = summary_border

    # Sum of module parent rows: 6, 9, 13, 17, 21, 25, 29, 33
    cell_dur = ws.cell(row=summary_row, column=7, value="=SUM(G6,G9,G13,G17,G21,G25,G29,G33)")
    cell_dur.font = font_bold_body
    cell_dur.fill = fill_summary
    cell_dur.alignment = align_right
    cell_dur.number_format = "0.0"
    cell_dur.border = summary_border

    cell_bobot = ws.cell(row=summary_row, column=8, value="=SUM(H6,H9,H13,H17,H21,H25,H29,H33)")
    cell_bobot.font = font_bold_body
    cell_bobot.fill = fill_summary
    cell_bobot.alignment = align_right
    cell_bobot.number_format = '0.0"%"'
    cell_bobot.border = summary_border

    ws.merge_cells(f"I{summary_row}:J{summary_row}")
    sum_stat = ws.cell(row=summary_row, column=9, value="Status: Siap Eksekusi (Ready)")
    sum_stat.font = font_bold_body
    sum_stat.fill = fill_summary
    sum_stat.alignment = align_center
    ws.cell(row=summary_row, column=10).border = summary_border
    sum_stat.border = summary_border

    output_path = r"C:\Users\Dell 3490\.gemini\antigravity\scratch\ciheul-helpdesk\docs\wbs\wbs_helpdesk_system.xlsx"
    wb.save(output_path)
    print(f"Base openpyxl saved: {output_path}")

    # ----------------------------------------------------
    # Native Excel Automation (Alt+H, O, I & Alt+H, O, A)
    # ----------------------------------------------------
    apply_excel_native_autofit(output_path)

def apply_excel_native_autofit(excel_file_path):
    import win32com.client
    try:
        excel = win32com.client.Dispatch("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        wb = excel.Workbooks.Open(os.path.abspath(excel_file_path))
        ws = wb.Sheets(1)

        # Ensure multi-line text wraps cleanly
        ws.Range("C6:D40").WrapText = True
        ws.Range("J6:J40").WrapText = True

        # Alt+H, O, I: AutoFit Columns
        ws.Columns.AutoFit()

        # Provide generous comfortable column widths so text has breathing room
        col_widths = {
            "A": 11,   # WBS ID
            "B": 24,   # Modul / Komponen
            "C": 28,   # Fitur / Work Package
            "D": 48,   # Rincian Task
            "E": 18,   # User Role
            "F": 24,   # Tech Stack
            "G": 14,   # Durasi
            "H": 12,   # Bobot
            "I": 14,   # Prioritas
            "J": 44    # Kriteria Verifikasi
        }
        for col_letter, min_w in col_widths.items():
            current_w = ws.Columns(col_letter).ColumnWidth
            if current_w < min_w:
                ws.Columns(col_letter).ColumnWidth = min_w

        # Alt+H, O, A: AutoFit Rows
        ws.Rows.AutoFit()

        # Adjust header rows cleanly
        ws.Rows(1).RowHeight = 28
        ws.Rows(2).RowHeight = 22
        ws.Rows(3).RowHeight = 20
        ws.Rows(4).RowHeight = 8
        ws.Rows(5).RowHeight = 26

        # Ensure all data rows have a comfortable minimum height (padding)
        for r in range(6, ws.UsedRange.Rows.Count + 1):
            rh = ws.Rows(r).RowHeight
            if rh < 24:
                ws.Rows(r).RowHeight = 24
            else:
                ws.Rows(r).RowHeight = rh + 4

        wb.Save()
        wb.Close(True)
        excel.Quit()
        print("Excel COM AutoFit (Alt+H, O, I and Alt+H, O, A) successfully applied!")
    except Exception as e:
        print(f"Warning during Excel COM automation: {e}")

if __name__ == "__main__":
    build_wbs_spreadsheet()
